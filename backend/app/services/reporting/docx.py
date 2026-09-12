"""DOCX report rendering — TRD FR-27, package B11.

The problem statement asks for an **editable** report, which is why this format exists at all: an
officer amends the narrative before filing, a brand's packaging team pastes findings into a
correction brief. So FR-27's acceptance test is specifically that the findings table opens in
Word and LibreOffice as a real table — a ``w:tbl`` — and not as a picture of one.

Renders from the same ``ReportData`` as the PDF and the JSON. It may choose how to present a
value; it may not choose which values exist.
"""

from __future__ import annotations

import io

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.services.reporting.model import ADVISORY_DISCLAIMER, ReportData

_VERDICT_COLOUR = {
    "FAIL": RGBColor(0xAA, 0x11, 0x11),
    "BORDERLINE": RGBColor(0xAA, 0x66, 0x00),
    "NOT_ASSESSABLE": RGBColor(0x55, 0x55, 0x55),
    "PASS": RGBColor(0x11, 0x66, 0x11),
}

_COLUMNS = ("Rule", "Verdict", "Observed", "Required", "Detail")


def _add_metadata(document: DocxDocument, data: ReportData) -> None:
    meta = data.metadata
    rows: list[tuple[str, str]] = [
        ("Scan ID", meta.scan_id),
        ("Product", meta.product_name or "—"),
        ("Organisation", meta.org_name),
        ("Mode", meta.mode),
        ("Captured", meta.captured_at.isoformat()),
        ("Report generated", meta.generated_at.isoformat()),
        ("Rule pack", meta.rulepack_version),
    ]
    if meta.inspector_name:
        rows.append(("Inspector", meta.inspector_name))
    if meta.location:
        rows.append(("Location", meta.location))
    rows.append(("Image SHA-256", meta.image_sha256))
    rows.append(("Findings SHA-256", meta.findings_sha256))

    for label, value in rows:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        run.font.size = Pt(9)
        value_run = paragraph.add_run(value)
        value_run.font.size = Pt(9)


def _add_findings_table(document: DocxDocument, data: ReportData) -> None:
    table = document.add_table(rows=1, cols=len(_COLUMNS))
    table.style = "Table Grid"

    for cell, heading in zip(table.rows[0].cells, _COLUMNS, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(heading)
        run.bold = True
        run.font.size = Pt(9)

    for finding in data.findings:
        cells = table.add_row().cells
        # Column 0 and 1 are what the drift test reads back: they must hold the machine-readable
        # rule id and verdict, not a prettified label, or the DOCX and the PDF stop comparable.
        cells[0].text = finding.rule_id
        cells[1].text = finding.verdict

        verdict_run = cells[1].paragraphs[0].runs[0]
        verdict_run.bold = True
        verdict_run.font.color.rgb = _VERDICT_COLOUR[finding.verdict]

        cells[2].text = finding.observed or "—"
        if finding.band:
            band_run = cells[2].paragraphs[0].add_run(f"\nmeasured band {finding.band}")
            band_run.italic = True
            band_run.font.size = Pt(7.5)

        cells[3].text = finding.required or "—"

        cells[4].text = finding.message
        citation_paragraph = cells[4].add_paragraph()
        citation_run = citation_paragraph.add_run(finding.citation)
        citation_run.italic = True
        citation_run.font.size = Pt(7.5)

        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is None:
                        run.font.size = Pt(8.5)


def render_docx_stream(data: ReportData) -> io.BytesIO:
    """Render the report into an in-memory DOCX stream, positioned at the start."""
    document = Document()

    heading = document.add_heading("Packaged commodity compliance report", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(
        "Legal Metrology (Packaged Commodities) Rules, 2011 · "
        f"rule pack {data.metadata.rulepack_version}"
    )
    subtitle_run.font.size = Pt(9)
    subtitle_run.italic = True

    document.add_heading("Scan details", level=2)
    _add_metadata(document, data)

    document.add_heading("Summary", level=2)
    summary = document.add_paragraph()
    for key, label in (
        ("fail", "Fail"),
        ("borderline", "Borderline"),
        ("na", "Not assessable"),
        ("pass", "Pass"),
        ("not_applicable", "Not applicable"),
    ):
        run = summary.add_run(f"{data.summary.get(key, 0)} {label}    ")
        run.font.size = Pt(9.5)

    document.add_heading("Findings", level=2)
    _add_findings_table(document, data)

    if data.not_applicable_rule_ids:
        document.add_heading("Rules that did not apply to this product", level=2)
        paragraph = document.add_paragraph()
        run = paragraph.add_run(
            "These rules were not evaluated because their conditions were not met by this "
            "product, or because they are not yet in force at the date of this scan. They are "
            "neither passes nor failures: "
            + ", ".join(data.not_applicable_rule_ids)
            + "."
        )
        run.font.size = Pt(8)

    if data.annotated_image_png is not None:
        document.add_heading("Annotated evidence", level=2)
        document.add_picture(io.BytesIO(data.annotated_image_png))

    document.add_paragraph()
    disclaimer = document.add_paragraph()
    disclaimer_run = disclaimer.add_run(ADVISORY_DISCLAIMER)
    disclaimer_run.font.size = Pt(8)
    disclaimer_run.italic = True

    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def render_docx(data: ReportData) -> bytes:
    """Render the report as DOCX bytes."""
    return render_docx_stream(data).getvalue()


__all__ = ["render_docx", "render_docx_stream"]
