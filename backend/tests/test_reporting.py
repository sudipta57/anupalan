"""Report generation — TRD FR-27, backend work package B11.

The acceptance test in the TRD is that the DOCX opens in Word and LibreOffice with the findings
table editable **as a real table, not an image**, and that the PDF carries the annotated image,
the findings table, the metadata and both hashes. Underneath that sits the requirement that
matters more: PDF and DOCX render from **one shared data structure**, so the two can never
disagree about what the verdict was.

That is what most of this file tests — not that each renderer works in isolation, but that they
agree. A PDF and a DOCX of the same scan showing different verdicts is the kind of defect that
surfaces in front of a magistrate, not in a test run.

**On the split between `render_html` and `render_pdf`.** WeasyPrint turns HTML into PDF and needs
the GTK3 native stack to do it. Content correctness does not depend on that step, so `pdf.py`
exposes the HTML separately and nearly every assertion below runs against it — on any machine,
with no native libraries. Only ``test_pdf_actually_rasterises`` needs WeasyPrint, and it skips
with a reason where the libraries are absent.
"""

from __future__ import annotations

import json
import re
from dataclasses import replace
from datetime import UTC, datetime

import pytest

from app.services.reporting import docx as docx_renderer
from app.services.reporting import json_report, model, pdf
from app.services.reporting.model import ADVISORY_DISCLAIMER, ReportData, ReportMetadata
from app.services.rules.evaluate import evaluate
from app.services.rules.findings import assemble
from app.services.rules.loader import active_pack
from app.services.rules.types import Extraction, Measurement, Profile

try:  # pragma: no cover - import probe, not logic
    import weasyprint  # noqa: F401

    WEASYPRINT_AVAILABLE = True
except Exception:  # noqa: BLE001 - any failure to load the native stack means the same thing
    WEASYPRINT_AVAILABLE = False

needs_weasyprint = pytest.mark.skipif(
    not WEASYPRINT_AVAILABLE,
    reason="WeasyPrint needs the GTK3 native stack (pango/cairo); absent on this machine",
)


# --------------------------------------------------------------------------- fixtures


@pytest.fixture(scope="module")
def report_data() -> ReportData:
    """One scan carried all the way through: evaluate -> assemble -> report model.

    Deliberately a label with a mix of verdicts, so the renderers are exercised on FAIL,
    BORDERLINE, NOT_ASSESSABLE and PASS rather than on a uniformly clean sheet.
    """
    pack = active_pack()
    profile = Profile(
        qty_basis="weight_or_volume",
        net_qty_in_g_or_ml=250.0,
        net_qty_value=250.0,
        net_qty_unit="g",
        surface="printed",
    )
    declarations = [
        Extraction(field_code="manufacturer_name", value_raw="Kalyani Foods Pvt Ltd"),
        Extraction(field_code="manufacturer_address", value_raw="12 GT Road, Kalyani 741235"),
        Extraction(field_code="common_name", value_raw="Roasted Chana"),
        Extraction(field_code="net_quantity", value_raw="250 gms"),
        Extraction(field_code="mfg_month_year", value_raw="03/2026"),
        Extraction(field_code="mrp", value_raw="₹250"),
        Extraction(field_code="consumer_care_name", value_raw="Consumer Care Cell"),
    ]
    measurements = [
        Measurement(
            field_code="net_quantity",
            glyph="2",
            height_mm=2.05,
            width_mm=1.3,
            uncertainty_mm=0.25,
            is_numeral=True,
            method="connected_components",
        )
    ]

    findings = assemble(
        evaluate(
            profile,
            declarations,
            measurements,
            rulepack=pack,
            as_of=datetime(2026, 10, 1, tzinfo=UTC).date(),
        ),
        pack,
    )

    return model.build(
        findings,
        ReportMetadata(
            scan_id="3f8c1d20-0000-4000-8000-000000000001",
            org_name="Nadia District Legal Metrology",
            product_name="Roasted Chana 250 g",
            mode="enforcement",
            captured_at=datetime(2026, 10, 1, 9, 30, tzinfo=UTC),
            generated_at=datetime(2026, 10, 1, 9, 31, tzinfo=UTC),
            rulepack_version=findings.rulepack_version,
            image_sha256="a" * 64,
            findings_sha256="b" * 64,
        ),
    )


def _html_verdicts(html: str) -> list[tuple[str, str]]:
    """Pull (rule_id, verdict) out of the rendered findings table."""
    return [
        (match.group("rule"), match.group("verdict"))
        for match in re.finditer(
            r'data-rule="(?P<rule>[^"]+)"\s+data-verdict="(?P<verdict>[^"]+)"', html
        )
    ]


def _docx_verdicts(document: object) -> list[tuple[str, str]]:
    """Pull (rule_id, verdict) out of the DOCX findings table, skipping its header row."""
    table = document.tables[0]  # type: ignore[attr-defined]
    return [
        (row.cells[0].text.strip(), row.cells[1].text.strip()) for row in table.rows[1:]
    ]


# --------------------------------------------------------------------------- the drift test


def test_pdf_and_docx_report_identical_findings(report_data: ReportData) -> None:
    """The requirement B11 exists to enforce: the two formats cannot disagree.

    Identical row count and identical verdict strings, in the same order. If this ever fails,
    one renderer has grown its own filtering or sorting and the shared data structure has stopped
    being shared.
    """
    import docx as docx_lib

    html = pdf.render_html(report_data)
    document = docx_lib.Document(docx_renderer.render_docx_stream(report_data))

    from_html = _html_verdicts(html)
    from_docx = _docx_verdicts(document)

    assert from_html, "the HTML rendered no findings rows at all"
    assert len(from_html) == len(from_docx)
    assert from_html == from_docx


def test_json_reports_the_same_findings_as_the_other_two(report_data: ReportData) -> None:
    payload = json.loads(json_report.render_json(report_data))
    html = _html_verdicts(pdf.render_html(report_data))

    assert [(f["rule_id"], f["verdict"]) for f in payload["findings"]] == html


# --------------------------------------------------------------------------- FR-27 specifics


def test_docx_findings_table_is_a_real_table_not_an_image(report_data: ReportData) -> None:
    """FR-27's acceptance test: editable in Word and LibreOffice, so it must be a w:tbl."""
    import docx as docx_lib

    document = docx_lib.Document(docx_renderer.render_docx_stream(report_data))

    assert document.tables, "the DOCX carries no table"
    assert document.tables[0]._tbl.tag.endswith("}tbl"), "findings table is not a real w:tbl"
    assert not document.inline_shapes or all(
        shape.type != 3 for shape in document.inline_shapes
    ), "findings must not be rendered as a picture"


def test_the_advisory_disclaimer_appears_in_every_format(report_data: ReportData) -> None:
    """CLAUDE.md §3.8 — this is a pre-audit tool, not a certification, and every artefact says so.

    Not configurable off. A report that omits it misrepresents what this product is.
    """
    import docx as docx_lib

    html = pdf.render_html(report_data)
    document = docx_lib.Document(docx_renderer.render_docx_stream(report_data))
    docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    payload = json.loads(json_report.render_json(report_data))

    assert ADVISORY_DISCLAIMER in html
    assert ADVISORY_DISCLAIMER in docx_text
    assert payload["disclaimer"] == ADVISORY_DISCLAIMER


def test_every_format_carries_the_rulepack_version(report_data: ReportData) -> None:
    """CLAUDE.md §3.6 — a report must name the rules it was issued under."""
    import docx as docx_lib

    version = report_data.metadata.rulepack_version
    assert version == "LM-2011-v1.0"

    document = docx_lib.Document(docx_renderer.render_docx_stream(report_data))
    docx_text = "\n".join(p.text for p in document.paragraphs) + "".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )

    assert version in pdf.render_html(report_data)
    assert version in docx_text
    payload = json.loads(json_report.render_json(report_data))
    assert payload["metadata"]["rulepack_version"] == version


def test_both_hashes_appear_in_the_report(report_data: ReportData) -> None:
    """Evidence integrity (architecture §10): the report embeds the image and findings hashes,
    so anyone can verify it was not altered after issue."""
    html = pdf.render_html(report_data)

    assert report_data.metadata.image_sha256 in html
    assert report_data.metadata.findings_sha256 in html


def test_rules_that_did_not_apply_are_listed(report_data: ReportData) -> None:
    """A reader must be able to tell which rules were checked.

    Silence about a skipped rule reads as "we checked and it was fine" (docs/decisions.md,
    2026-09-12).
    """
    html = pdf.render_html(report_data)

    assert report_data.not_applicable_rule_ids
    for rule_id in report_data.not_applicable_rule_ids:
        assert rule_id in html


def test_borderline_bands_are_printed_not_rounded_away(report_data: ReportData) -> None:
    """A BORDERLINE verdict is only defensible if the band that produced it is visible."""
    borderline = [f for f in report_data.findings if f.verdict == "BORDERLINE"]
    assert borderline, "fixture should produce at least one BORDERLINE finding"

    html = pdf.render_html(report_data)
    for finding in borderline:
        assert finding.band is not None
        assert finding.band in html


def test_html_escapes_values_that_came_off_a_label(report_data: ReportData) -> None:
    """Extracted text is untrusted input. It came off a photograph of a package, went through
    OCR and possibly an LLM, and lands in a document somebody opens. It must not be able to
    inject markup."""
    hostile = replace(
        report_data,
        metadata=replace(report_data.metadata, product_name="<script>alert(1)</script>"),
    )

    html = pdf.render_html(hostile)

    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html


# --------------------------------------------------------------------------- determinism


def test_json_output_is_byte_stable(report_data: ReportData) -> None:
    """Two renders of one report must be identical, or the hash chain over it is meaningless."""
    assert json_report.render_json(report_data) == json_report.render_json(report_data)


# --------------------------------------------------------------------------- the native step


@needs_weasyprint
def test_pdf_actually_rasterises(report_data: ReportData) -> None:
    """The one assertion that needs the GTK3 stack. Everything else above runs anywhere."""
    blob = pdf.render_pdf(report_data)

    assert blob.startswith(b"%PDF-")
    assert len(blob) > 1000
